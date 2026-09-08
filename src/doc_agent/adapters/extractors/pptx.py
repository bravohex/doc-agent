# pyright: reportUnknownVariableType=false, reportUnknownMemberType=false, reportUnknownArgumentType=false, reportUnknownParameterType=false, reportMissingParameterType=false, reportMissingTypeArgument=false, reportAttributeAccessIssue=false, reportPrivateUsage=false, reportCallIssue=false, reportArgumentType=false
# Dynamic third-party Office/UI APIs expose incomplete static types; strict checking remains enabled for domain, application, ports, and typed adapters.
"""PowerPoint extraction preserving slide/shape identity and spatial metadata."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from doc_agent.adapters.extractors.failures import readable
from doc_agent.adapters.extractors.ooxml import SafeOoxmlPackage
from doc_agent.adapters.extractors.properties import core_properties, withheld
from doc_agent.domain.identifiers import stable_key
from doc_agent.domain.models import (
    Block,
    BlockKind,
    Container,
    ExtractedDocument,
    ExtractedVisual,
    PptxLocator,
)


class PptxExtractor:
    """Extract slide text/tables/notes while preserving visual-inspection signals."""

    name = "pptx"
    version = "1.0"

    def supports(self, source: Path) -> bool:
        return source.suffix.lower() == ".pptx"

    def extract(self, source: Path) -> ExtractedDocument:
        SafeOoxmlPackage(source).inspect()
        with readable(source, "PPTX presentation"):
            prs = Presentation(source)
        containers: list[Container] = []
        blocks: list[Block] = []
        visuals: list[ExtractedVisual] = []

        for slide_number, slide in enumerate(prs.slides, start=1):
            title = (
                slide.shapes.title.text.strip()
                if slide.shapes.title and slide.shapes.title.text
                else f"Slide {slide_number}"
            )
            visual_required = self._visual_required(slide.shapes)
            container_key = stable_key("pptx", "slide", title, hint=slide_number)
            containers.append(
                Container(
                    stable_key=container_key,
                    kind="slide",
                    title=title,
                    ordinal=slide_number,
                    source=PptxLocator(slide_number=slide_number),
                    metadata={
                        "shape_count": len(slide.shapes),
                        # A slide set never to show is the deck's hidden worksheet: it
                        # travels with the file and is skipped when presented.
                        "hidden": self._hidden(slide),
                        "layout": self._layout_name(slide),
                        "has_notes": bool(slide.has_notes_slide),
                    },
                    visual_required=visual_required,
                )
            )
            ordinal = 0
            for shape in slide.shapes:
                ordinal += 1
                locator = PptxLocator(
                    slide_number=slide_number,
                    shape_id=shape.shape_id,
                    shape_name=shape.name,
                )
                presentation = self._geometry(shape)
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text:
                        blocks.append(
                            Block(
                                stable_key=stable_key(
                                    "pptx", title, f"shape:{shape.shape_id}", hint=text[:80]
                                ),
                                container_key=container_key,
                                kind=BlockKind.TEXT_BOX,
                                ordinal=ordinal,
                                text=text,
                                source=locator,
                                presentation=presentation,
                                visual_required=visual_required,
                            )
                        )
                if getattr(shape, "has_table", False):
                    for row_index, row in enumerate(shape.table.rows, start=1):
                        cells = [cell.text.strip() for cell in row.cells]
                        text = "\t".join(value for value in cells if value)
                        if not text:
                            continue
                        blocks.append(
                            Block(
                                stable_key=stable_key(
                                    "pptx",
                                    title,
                                    cells[0] or text[:80],
                                    hint=f"{shape.shape_id}:{row_index}",
                                ),
                                container_key=container_key,
                                kind=BlockKind.TABLE_ROW,
                                ordinal=ordinal * 1000 + row_index,
                                text=text,
                                source=PptxLocator(
                                    slide_number=slide_number,
                                    shape_id=shape.shape_id,
                                    shape_name=shape.name,
                                    row_index=row_index,
                                ),
                                payload={"cells": cells},
                                presentation=presentation,
                            )
                        )
                if getattr(shape, "has_chart", False):
                    chart_payload = self._chart_payload(shape.chart)
                    blocks.append(
                        Block(
                            stable_key=stable_key("pptx", title, f"chart:{shape.shape_id}"),
                            container_key=container_key,
                            kind=BlockKind.CHART,
                            ordinal=ordinal,
                            text=chart_payload.get("title") or shape.name,
                            source=locator,
                            payload=chart_payload,
                            presentation=presentation,
                            visual_required=True,
                        )
                    )
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image = shape.image
                    visuals.append(
                        ExtractedVisual(
                            stable_key=stable_key("pptx", title, f"image:{shape.shape_id}"),
                            media_type=image.content_type,
                            source=locator,
                            data=image.blob,
                            width=self._emu(shape.width),
                            height=self._emu(shape.height),
                            alt_text=shape.name,
                        )
                    )

            try:
                notes_frame = slide.notes_slide.notes_text_frame
                notes_text = notes_frame.text.strip() if notes_frame is not None else ""
            except Exception:
                notes_text = ""
            if notes_text:
                blocks.append(
                    Block(
                        stable_key=stable_key("pptx", title, "notes", hint=notes_text[:80]),
                        container_key=container_key,
                        kind=BlockKind.NOTE,
                        ordinal=99_999,
                        text=notes_text,
                        source=PptxLocator(slide_number=slide_number, shape_name="notes"),
                    )
                )

        hidden_slides = [
            container.ordinal for container in containers if container.metadata.get("hidden")
        ]
        return ExtractedDocument(
            logical_name=source.name,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            source_path=str(source),
            containers=containers,
            blocks=blocks,
            visuals=visuals,
            metadata={
                "properties": core_properties(prs.core_properties),
                "slide_count": len(containers),
                "hidden_slides": hidden_slides,
                "slide_size": {
                    "width_emu": prs.slide_width,
                    "height_emu": prs.slide_height,
                },
                "withheld_content": withheld(
                    f"Slide(s) {', '.join(str(number) for number in hidden_slides)} are "
                    "hidden, so they are extracted here but skipped when the deck is shown."
                    if hidden_slides
                    else None,
                ),
            },
        )

    @staticmethod
    def _hidden(slide: Any) -> bool:
        """Read the slide's show flag.

        python-pptx does not model it, so the attribute is read directly; absent means
        shown, which is the format's default.
        """

        return slide._element.get("show") in ("0", "false")

    @staticmethod
    def _layout_name(slide: Any) -> str | None:
        try:
            return slide.slide_layout.name
        except Exception:
            return None

    @staticmethod
    def _emu(value: Any) -> int | None:
        """Shapes inheriting layout geometry report ``None`` for position and size."""

        return int(value) if value is not None else None

    @classmethod
    def _geometry(cls, shape: Any) -> dict[str, Any]:
        return {
            "left": cls._emu(shape.left),
            "top": cls._emu(shape.top),
            "width": cls._emu(shape.width),
            "height": cls._emu(shape.height),
            "shape_type": str(shape.shape_type),
        }

    @staticmethod
    def _visual_required(shapes: Any) -> bool:
        spatial = 0
        for shape in shapes:
            if shape.shape_type in {
                MSO_SHAPE_TYPE.GROUP,
                MSO_SHAPE_TYPE.AUTO_SHAPE,
                MSO_SHAPE_TYPE.FREEFORM,
            }:
                spatial += 1
        return spatial >= 2 or len(shapes) >= 8

    @staticmethod
    def _chart_payload(chart: Any) -> dict[str, Any]:
        title = ""
        try:
            title = chart.chart_title.text_frame.text.strip() if chart.has_title else ""
        except Exception:
            title = ""
        series: list[dict[str, Any]] = []
        for item in chart.series:
            try:
                values = list(item.values)
            except Exception:
                values = []
            series.append({"name": str(item.name or ""), "values": values})
        return {"title": title, "series": series}
