# pyright: reportUnknownVariableType=false, reportUnknownMemberType=false, reportUnknownArgumentType=false, reportUnknownParameterType=false, reportMissingParameterType=false, reportMissingTypeArgument=false, reportAttributeAccessIssue=false, reportPrivateUsage=false, reportCallIssue=false, reportArgumentType=false
# Dynamic third-party Office/UI APIs expose incomplete static types; strict checking remains enabled for domain, application, ports, and typed adapters.
"""DOCX extraction preserving document order and heading hierarchy."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from doc_agent.adapters.extractors.failures import readable
from doc_agent.adapters.extractors.ooxml import SafeOoxmlPackage
from doc_agent.adapters.extractors.properties import core_properties, withheld
from doc_agent.domain.identifiers import stable_key
from doc_agent.domain.models import (
    Block,
    BlockKind,
    Container,
    DocxLocator,
    ExtractedDocument,
    ExtractedVisual,
)


class DocxExtractor:
    """Extract semantic text and tables without relying on unstable rendered page numbers."""

    name = "docx"
    version = "1.0"

    def supports(self, source: Path) -> bool:
        return source.suffix.lower() == ".docx"

    def extract(self, source: Path) -> ExtractedDocument:
        SafeOoxmlPackage(source).inspect()
        with readable(source, "DOCX document"):
            document = Document(source)
        blocks: list[Block] = []
        containers: list[Container] = []
        visuals: list[ExtractedVisual] = []
        headings: list[str] = []
        section_key = stable_key("docx", "document", "root")
        containers.append(
            Container(
                stable_key=section_key,
                kind="section",
                title="Document",
                ordinal=1,
                source=DocxLocator(),
            )
        )
        paragraph_index = 0
        table_index = 0
        ordinal = 0
        for item in self._iter_block_items(document):
            ordinal += 1
            if isinstance(item, Paragraph):
                paragraph_index += 1
                text = item.text.strip()
                if not text:
                    continue
                style = (item.style.name or "") if item.style else ""
                level = self._heading_level(style)
                kind = BlockKind.PARAGRAPH
                if level is not None:
                    headings = headings[: level - 1]
                    headings.append(text)
                    kind = BlockKind.HEADING
                    section_key = stable_key("docx", "heading", "/".join(headings))
                    containers.append(
                        Container(
                            stable_key=section_key,
                            kind="section",
                            title=text,
                            ordinal=len(containers) + 1,
                            source=DocxLocator(
                                section_path=tuple(headings), paragraph_index=paragraph_index
                            ),
                        )
                    )
                elif style.lower().startswith("list"):
                    kind = BlockKind.LIST_ITEM
                locator = DocxLocator(section_path=tuple(headings), paragraph_index=paragraph_index)
                blocks.append(
                    Block(
                        stable_key=stable_key(
                            "docx", "/".join(headings) or "root", text, hint=style
                        ),
                        container_key=section_key,
                        kind=kind,
                        ordinal=ordinal,
                        text=text,
                        source=locator,
                        payload={"hyperlinks": self._hyperlinks(item)},
                        presentation={"style": style},
                    )
                )
            else:
                table_index += 1
                for row_index, row in enumerate(item.rows, start=1):
                    texts = [cell.text.strip() for cell in row.cells]
                    text = "\t".join(value for value in texts if value)
                    if not text:
                        continue
                    locator = DocxLocator(
                        section_path=tuple(headings), table_index=table_index, row_index=row_index
                    )
                    identity = texts[0] if texts and texts[0] else text[:80]
                    blocks.append(
                        Block(
                            stable_key=stable_key(
                                "docx",
                                "/".join(headings) or "root",
                                identity,
                                hint=f"table:{table_index}",
                            ),
                            container_key=section_key,
                            kind=BlockKind.TABLE_ROW,
                            ordinal=ordinal * 1000 + row_index,
                            text=text,
                            source=locator,
                            payload={"cells": texts},
                            presentation={"table_index": table_index},
                        )
                    )

        for section_no, section in enumerate(document.sections, start=1):
            for kind, paragraphs in (
                (BlockKind.HEADER, section.header.paragraphs),
                (BlockKind.FOOTER, section.footer.paragraphs),
            ):
                for idx, paragraph in enumerate(paragraphs, start=1):
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    blocks.append(
                        Block(
                            stable_key=stable_key("docx", kind.value, text, hint=section_no),
                            container_key=stable_key("docx", kind.value, section_no),
                            kind=kind,
                            ordinal=100_000 + section_no * 100 + idx,
                            text=text,
                            source=DocxLocator(part=kind.value, paragraph_index=idx),
                        )
                    )

        for rel_id, rel in document.part.rels.items():
            if "image" not in rel.reltype:
                continue
            part = rel.target_part
            blob = part.blob
            content_type = getattr(part, "content_type", "application/octet-stream")
            visuals.append(
                ExtractedVisual(
                    stable_key=stable_key("docx", "media", rel_id, hint=part.partname),
                    media_type=content_type,
                    source=DocxLocator(part=str(part.partname)),
                    data=blob,
                )
            )

        revisions = self._revisions(document)
        hidden_runs = len(self._find(document, "vanish"))
        fields = self._fields(document)
        return ExtractedDocument(
            logical_name=source.name,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            source_path=str(source),
            containers=containers,
            blocks=blocks,
            visuals=visuals,
            metadata={
                "properties": core_properties(document.core_properties),
                "protection": self._protection(document),
                "revisions": revisions,
                "hidden_text_runs": hidden_runs,
                "fields": fields,
                "withheld_content": withheld(
                    f"{revisions['deletions']} tracked deletion(s) are still in the file: "
                    "text shown as removed has not been accepted and remains present."
                    if revisions["deletions"]
                    else None,
                    f"{revisions['insertions']} tracked insertion(s) are unaccepted, so the "
                    "text read here is not the document as last agreed."
                    if revisions["insertions"]
                    else None,
                    f"{hidden_runs} run(s) are marked hidden and do not print or display."
                    if hidden_runs
                    else None,
                ),
            },
        )

    @staticmethod
    def _find(document: DocumentObject, tag: str) -> list[Any]:
        """Find every ``w:<tag>`` in the body.

        python-docx models paragraphs and tables, not revision marks, so these are read
        from the XML directly rather than inferred from the text.
        """

        return list(document.element.body.iter(qn(f"w:{tag}")))

    @classmethod
    def _revisions(cls, document: DocumentObject) -> dict[str, Any]:
        """Count tracked changes, and recover the text a deletion still carries.

        This is the finding with no spreadsheet counterpart and the largest consequence:
        a document with unaccepted revisions is not the document it appears to be, and
        deleted text remains in the file rather than being gone from it.
        """

        insertions = cls._find(document, "ins")
        deletions = cls._find(document, "del")
        authors = sorted(
            {
                str(element.get(qn("w:author")))
                for element in (*insertions, *deletions)
                if element.get(qn("w:author"))
            }
        )
        deleted_text = [element.text for element in cls._find(document, "delText") if element.text]
        return {
            "insertions": len(insertions),
            "deletions": len(deletions),
            "authors": authors,
            "deleted_text": deleted_text,
        }

    @classmethod
    def _fields(cls, document: DocumentObject) -> list[dict[str, Any]]:
        """Record field codes with the result stored for them.

        A field is the document counterpart of a spreadsheet formula: the text on the
        page is a saved result, and nothing here recalculates it. The instruction is
        kept beside the result so a stale date or cross-reference is visible as such.
        """

        recorded: list[dict[str, Any]] = []
        for element in cls._find(document, "fldSimple"):
            instruction = str(element.get(qn("w:instr")) or "").strip()
            result = "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()
            recorded.append(
                {
                    "instruction": instruction,
                    "result": result,
                    # Same vocabulary as a spreadsheet cell: a result that was saved, or
                    # a field the file carries no result for.
                    "value_state": "cached" if result else "uncalculated",
                }
            )
        return recorded

    @staticmethod
    def _protection(document: DocumentObject) -> dict[str, Any]:
        """Record an editing restriction, the document counterpart of sheet protection."""

        try:
            settings = document.settings.element
        except Exception:
            return {"enabled": False}
        for element in settings.iter(qn("w:documentProtection")):
            return {
                "enabled": True,
                "edit": element.get(qn("w:edit")),
                "enforced": element.get(qn("w:enforcement")) in ("1", "true", "on"),
            }
        return {"enabled": False}

    @staticmethod
    def _iter_block_items(document: DocumentObject) -> Iterator[Paragraph | Table]:
        for child in document.element.body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, document)
            elif isinstance(child, CT_Tbl):
                yield Table(child, document)

    @staticmethod
    def _heading_level(style_name: str) -> int | None:
        if style_name.lower().startswith("heading "):
            try:
                return int(style_name.rsplit(" ", 1)[1])
            except ValueError:
                return None
        return None

    @staticmethod
    def _hyperlinks(paragraph: Paragraph) -> list[dict[str, str]]:
        result: list[dict[str, str]] = []
        namespaces = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        }
        for link in paragraph._p.xpath("./w:hyperlink"):
            rel_id = link.get(f"{{{namespaces['r']}}}id")
            text = "".join(node.text or "" for node in link.xpath(".//w:t"))
            if rel_id and rel_id in paragraph.part.rels:
                result.append({"text": text, "target": paragraph.part.rels[rel_id].target_ref})
        return result
