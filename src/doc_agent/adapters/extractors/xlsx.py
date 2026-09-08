# pyright: reportUnknownVariableType=false, reportUnknownMemberType=false, reportUnknownArgumentType=false, reportUnknownParameterType=false, reportMissingParameterType=false, reportMissingTypeArgument=false, reportAttributeAccessIssue=false, reportPrivateUsage=false, reportCallIssue=false, reportArgumentType=false
# Dynamic third-party Office/UI APIs expose incomplete static types; strict checking remains enabled for domain, application, ports, and typed adapters.
"""High-fidelity XLSX extraction into compact row blocks plus sparse cell metadata."""

from __future__ import annotations

import datetime as dt
import mimetypes
from pathlib import Path
from typing import Any, ClassVar

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from doc_agent.adapters.extractors.failures import readable
from doc_agent.adapters.extractors.ooxml import SafeOoxmlPackage
from doc_agent.adapters.extractors.properties import core_properties, withheld
from doc_agent.domain.identifiers import normalize_identity, stable_key
from doc_agent.domain.models import (
    Block,
    BlockKind,
    Container,
    ExtractedDocument,
    ExtractedVisual,
    XlsxLocator,
)

#: The automatic background and text colours, carried by cells that were never coloured.
_AUTOMATIC_THEMES = frozenset({0, 1})


def _as_ranges(numbers: list[int]) -> list[list[int]]:
    """Compress sorted numbers into inclusive ranges: [3, 4, 5, 9] -> [[3, 5], [9, 9]]."""

    ranges: list[list[int]] = []
    for number in numbers:
        if ranges and number == ranges[-1][1] + 1:
            ranges[-1][1] = number
            continue
        ranges.append([number, number])
    return ranges


class XlsxExtractor:
    """Extract spreadsheet values without collapsing raw, cached, and display representations."""

    name = "xlsx"
    version = "1.0"
    suffixes: ClassVar[frozenset[str]] = frozenset({".xlsx", ".xlsm", ".xltx", ".xltm"})

    def supports(self, source: Path) -> bool:
        return source.suffix.lower() in self.suffixes

    def extract(self, source: Path) -> ExtractedDocument:
        SafeOoxmlPackage(source).inspect()
        with readable(source, "XLSX workbook"):
            formula_wb = load_workbook(source, data_only=False, read_only=False)
            cached_wb = load_workbook(source, data_only=True, read_only=False)
        blocks: list[Block] = []
        containers: list[Container] = []
        visuals: list[ExtractedVisual] = []

        for sheet_no, ws in enumerate(formula_wb.worksheets, start=1):
            cached_ws = cached_wb[ws.title]
            container_key = stable_key("xlsx", "workbook", ws.title)
            locator = XlsxLocator(sheet=ws.title, row=1, cell_range=ws.dimensions)
            tables = [{"name": table.name, "ref": table.ref} for table in ws.tables.values()]
            containers.append(
                Container(
                    stable_key=container_key,
                    kind="worksheet",
                    title=ws.title,
                    ordinal=sheet_no,
                    source=locator,
                    metadata={
                        "state": ws.sheet_state,
                        "dimension": ws.dimensions,
                        "tables": tables,
                        # Rules govern a range rather than a cell, so they belong to the
                        # sheet: attaching them per cell would repeat one rule hundreds
                        # of times and still lose the range it applies to.
                        "validations": self._validations(ws),
                        "conditional_formats": self._conditional_formats(ws),
                        "defined_names": self._defined_names(ws.defined_names, scope=ws.title),
                        "layout": self._sheet_layout(ws),
                    },
                )
            )
            merged_lookup: dict[str, str] = {}
            for merged in ws.merged_cells.ranges:
                for row in ws.iter_rows(
                    min_row=merged.min_row,
                    max_row=merged.max_row,
                    min_col=merged.min_col,
                    max_col=merged.max_col,
                ):
                    for cell in row:
                        merged_lookup[cell.coordinate] = merged.coord

            # Row coordinates are intentionally excluded from stable identity. A row that is
            # physically moved by inserting content above it must remain the same semantic block.
            # For duplicate first-column identities we keep a deterministic occurrence number;
            # truly indistinguishable duplicate rows cannot be matched more precisely without a
            # durable business identifier supplied by the workbook itself.
            row_identity_occurrences: dict[str, int] = {}
            for row_no in range(1, ws.max_row + 1):
                # Semantics and position/style are kept apart so that moving or restyling a
                # row is not mistaken for a change in what the row says. The two lists share
                # one index, and each presentation entry names its own cell coordinate.
                row_cells: list[dict[str, Any]] = []
                row_layout: list[dict[str, Any]] = []
                display_values: list[str] = []
                first_identity: str | None = None
                for col_no in range(1, ws.max_column + 1):
                    cell = ws.cell(row_no, col_no)
                    if isinstance(cell, MergedCell):
                        continue
                    cached = cached_ws.cell(row_no, col_no).value
                    has_metadata = bool(
                        cell.comment or cell.hyperlink or cell.coordinate in merged_lookup
                    )
                    if cell.value is None and cached is None and not has_metadata:
                        continue
                    formula = cell.value if cell.data_type == "f" else None
                    raw_value = None if formula is not None else self._json_value(cell.value)
                    display = self._display(
                        cached if formula is not None else cell.value, cell.number_format
                    )
                    if display:
                        display_values.append(display)
                    if first_identity is None and cell.value not in (None, ""):
                        first_identity = normalize_identity(cell.value)
                    row_cells.append(
                        {
                            "raw_value": raw_value,
                            "formula": formula,
                            "cached_value": self._json_value(cached)
                            if formula is not None
                            else None,
                            "data_type": cell.data_type,
                            "display": display,
                            "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
                            "comment": cell.comment.text if cell.comment else None,
                        }
                    )
                    row_layout.append(
                        {
                            "coordinate": cell.coordinate,
                            "number_format": cell.number_format,
                            "merged_range": merged_lookup.get(cell.coordinate),
                            "hidden_column": bool(
                                ws.column_dimensions[get_column_letter(col_no)].hidden
                            ),
                            # Sparse by construction: a plain cell adds nothing here.
                            **self._cell_style(cell),
                        }
                    )
                if not row_cells:
                    continue
                identity = first_identity or f"content:{'|'.join(display_values)[:120]}"
                occurrence = row_identity_occurrences.get(identity, 0) + 1
                row_identity_occurrences[identity] = occurrence
                row_locator = XlsxLocator(
                    sheet=ws.title,
                    row=row_no,
                    cell_range=f"A{row_no}:{get_column_letter(ws.max_column)}{row_no}",
                )
                blocks.append(
                    Block(
                        stable_key=stable_key("xlsx", ws.title, identity, hint=occurrence),
                        container_key=container_key,
                        kind=BlockKind.TABLE_ROW,
                        ordinal=row_no,
                        text="\t".join(display_values),
                        source=row_locator,
                        payload={"cells": row_cells},
                        presentation={
                            "hidden_row": bool(ws.row_dimensions[row_no].hidden),
                            "cells": row_layout,
                        },
                    )
                )

            for image_no, image in enumerate(getattr(ws, "_images", []), start=1):
                try:
                    data = image._data()  # openpyxl exposes no public binary accessor.
                except Exception:
                    continue
                media_type = (
                    mimetypes.guess_type(getattr(image, "path", "image.png"))[0] or "image/png"
                )
                anchor_row = getattr(getattr(image, "anchor", None), "_from", None)
                row = (anchor_row.row + 1) if anchor_row is not None else None
                visuals.append(
                    ExtractedVisual(
                        stable_key=stable_key(
                            "xlsx", ws.title, f"image:{image_no}", hint=row or ""
                        ),
                        media_type=media_type,
                        source=XlsxLocator(sheet=ws.title, row=row),
                        data=data,
                        width=int(image.width) if image.width else None,
                        height=int(image.height) if image.height else None,
                    )
                )

            for chart_no, chart in enumerate(getattr(ws, "_charts", []), start=1):
                title = self._chart_title(chart) or f"Chart {chart_no}"
                blocks.append(
                    Block(
                        stable_key=stable_key("xlsx", ws.title, f"chart:{chart_no}", hint=title),
                        container_key=container_key,
                        kind=BlockKind.CHART,
                        ordinal=ws.max_row + chart_no,
                        text=title,
                        source=XlsxLocator(sheet=ws.title),
                        payload={"series_count": len(getattr(chart, "ser", []))},
                        visual_required=True,
                    )
                )

        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        hidden_sheets = [
            str(container.title)
            for container in containers
            if container.metadata.get("state") != "visible"
        ]
        return ExtractedDocument(
            logical_name=source.name,
            media_type=media_type,
            source_path=str(source),
            containers=containers,
            blocks=blocks,
            visuals=visuals,
            metadata={
                "properties": core_properties(formula_wb.properties),
                "sheet_count": len(formula_wb.worksheets),
                "calculation": self._calculation(formula_wb),
                "defined_names": self._defined_names(formula_wb.defined_names, scope="workbook"),
                "withheld_content": withheld(
                    f"Sheet(s) {', '.join(hidden_sheets)} are hidden, so they are extracted "
                    "and searchable here but are not visible when the workbook is opened."
                    if hidden_sheets
                    else None,
                ),
            },
        )

    @staticmethod
    def _colour(colour: Any) -> str | None:
        """Read an ARGB colour, refusing openpyxl's placeholder for "not this kind".

        ``Color.rgb`` and ``Color.theme`` return a validation *message* as their value
        when the colour is of the other kind, so a theme-coloured cell yields the string
        ``"Values must be of type <class 'str'>"`` for ``rgb``. Storing that would put a
        sentence where a colour belongs, so only a real ARGB value is accepted and a
        theme colour is reported as a theme instead.
        """

        if colour is None:
            return None
        rgb = getattr(colour, "rgb", None)
        if isinstance(rgb, str) and len(rgb) == 8:
            try:
                int(rgb, 16)
            except ValueError:
                return None
            return rgb
        theme = getattr(colour, "theme", None)
        if not isinstance(theme, int) or theme in _AUTOMATIC_THEMES:
            # Themes 0 and 1 are the automatic background and text colours that nearly
            # every cell carries, so recording them would make the sparse style map
            # dense again while saying nothing a reviewer could act on.
            return None
        # A theme index is reported as an index: the workbook's palette is not resolved
        # here, so naming a concrete colour would be a guess.
        return f"theme:{theme}"

    @classmethod
    def _cell_style(cls, cell: Any) -> dict[str, Any]:
        """Record only the styling that differs from a plain cell.

        Storing every font attribute of every cell would multiply the size of a workbook
        for facts that are almost always the default. What is kept is what a reviewer
        reads meaning into: emphasis, a struck-through row, a colour used as a status,
        and a cell left editable on a protected sheet.
        """

        style: dict[str, Any] = {}
        font = getattr(cell, "font", None)
        if font is not None:
            if font.bold:
                style["bold"] = True
            if font.italic:
                style["italic"] = True
            if font.strike:
                style["strikethrough"] = True
            colour = cls._colour(font.color)
            if colour:
                style["font_color"] = colour
        fill = getattr(cell, "fill", None)
        if fill is not None and fill.fill_type:
            colour = cls._colour(fill.start_color)
            if colour:
                style["fill_color"] = colour
        protection = getattr(cell, "protection", None)
        if protection is not None and protection.locked is False:
            # Only the exception is worth recording: cells are locked by default, and an
            # unlocked one is editable wherever the sheet is protected.
            style["locked"] = False
        return style

    @classmethod
    def _sheet_layout(cls, ws: Any) -> dict[str, Any]:
        """Record sheet-wide layout that hides or frames content.

        Each of these can conceal data from a reader: a hidden column, a zero width, a
        filter, rows folded away. They are sheet-wide, so they cost one record per sheet
        rather than one per cell.
        """

        columns: list[dict[str, Any]] = []
        for letter, dimension in sorted(ws.column_dimensions.items()):
            if not (dimension.hidden or dimension.customWidth):
                continue
            columns.append(
                {
                    "column": letter,
                    "width": dimension.width,
                    "hidden": bool(dimension.hidden),
                }
            )
        hidden_rows = sorted(
            int(number) for number, dimension in ws.row_dimensions.items() if dimension.hidden
        )
        return {
            "freeze_panes": ws.freeze_panes,
            "auto_filter": ws.auto_filter.ref if ws.auto_filter else None,
            "protected": bool(ws.protection.sheet),
            "columns": columns,
            # Ranges rather than every number, so a sheet with thousands of folded rows
            # stays a short answer.
            "hidden_rows": _as_ranges(hidden_rows),
        }

    @staticmethod
    def _validations(ws: Any) -> list[dict[str, Any]]:
        """Record the input rules a sheet enforces, with the ranges they cover.

        An audit asks what a cell was allowed to contain, which the value alone cannot
        answer: a status column restricted to a list is a different fact from a free
        text column that happens to hold the same word.
        """

        recorded: list[dict[str, Any]] = []
        for rule in ws.data_validations.dataValidation:
            recorded.append(
                {
                    "type": rule.type,
                    "operator": rule.operator,
                    "formula1": rule.formula1,
                    "formula2": rule.formula2,
                    "ranges": [str(part) for part in rule.sqref.ranges] if rule.sqref else [],
                    "allow_blank": bool(rule.allow_blank),
                    "show_error_message": bool(rule.showErrorMessage),
                    "error_title": rule.errorTitle,
                    "error_message": rule.error,
                    "prompt_title": rule.promptTitle,
                    "prompt_message": rule.prompt,
                }
            )
        return recorded

    @staticmethod
    def _conditional_formats(ws: Any) -> list[dict[str, Any]]:
        """Record conditional rules as rules, not as the appearance they produce.

        The colour a rule paints is not extracted: what matters for review is the
        condition and the range it is tested over.
        """

        recorded: list[dict[str, Any]] = []
        for group in ws.conditional_formatting:
            ranges = [str(part) for part in group.sqref.ranges] if group.sqref else []
            for rule in group.rules:
                recorded.append(
                    {
                        "ranges": ranges,
                        "type": rule.type,
                        "operator": rule.operator,
                        "formula": list(rule.formula) if rule.formula else [],
                        "priority": rule.priority,
                        "stop_if_true": bool(rule.stopIfTrue),
                    }
                )
        return recorded

    @staticmethod
    def _defined_names(names: Any, *, scope: str) -> list[dict[str, Any]]:
        """Record named ranges, which formulas reference instead of addresses."""

        recorded: list[dict[str, Any]] = []
        for name, definition in names.items():
            recorded.append(
                {
                    "name": name,
                    "refers_to": definition.attr_text,
                    "scope": scope,
                    "comment": definition.comment,
                    "hidden": bool(definition.hidden),
                }
            )
        return recorded

    @staticmethod
    def _calculation(workbook: Any) -> dict[str, Any]:
        """Record how the workbook was set to calculate.

        ``manual`` is the finding that matters: cached formula results in such a file
        may be stale by the workbook's own design, not by accident.
        """

        properties = workbook.calculation
        mode = getattr(properties, "calcMode", None) or "auto"
        return {
            "mode": mode,
            "automatic": mode != "manual",
            "full_calc_on_load": bool(getattr(properties, "fullCalcOnLoad", False)),
            "iterative": bool(getattr(properties, "iterate", False)),
            "iterate_count": getattr(properties, "iterateCount", None),
        }

    @staticmethod
    def _json_value(value: Any) -> Any:
        if isinstance(value, (dt.datetime, dt.date, dt.time)):
            return value.isoformat()
        return value

    @staticmethod
    def _display(value: Any, number_format: str) -> str:
        if value is None:
            return ""
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, (dt.datetime, dt.date, dt.time)):
            return value.isoformat()
        if isinstance(value, (int, float)) and "%" in number_format:
            decimals = 0
            if "." in number_format:
                decimals = len(number_format.split(".", 1)[1].split("%", 1)[0])
            return f"{value * 100:.{decimals}f}%"
        return str(value)

    @staticmethod
    def _chart_title(chart: Any) -> str | None:
        try:
            rich = chart.title.tx.rich
            return "".join(run.t for p in rich.p for run in p.r if run.t).strip() or None
        except Exception:
            return None
