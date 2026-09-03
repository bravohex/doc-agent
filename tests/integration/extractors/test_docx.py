from __future__ import annotations

from pathlib import Path

from docx import Document

from doc_agent.adapters.extractors.docx import DocxExtractor
from doc_agent.domain.models import BlockKind


def test_docx_preserves_heading_hierarchy_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Migration", level=1)
    doc.add_paragraph("Run two rehearsals before cutover.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "ID"
    table.cell(0, 1).text = "Requirement"
    table.cell(1, 0).text = "REQ-1"
    table.cell(1, 1).text = "Rollback"
    doc.sections[0].header.paragraphs[0].text = "Confidential"
    doc.save(path)

    extracted = DocxExtractor().extract(path)
    headings = [b for b in extracted.blocks if b.kind == BlockKind.HEADING]
    assert headings[0].text == "Migration"
    paragraph = next(b for b in extracted.blocks if "rehearsals" in b.text)
    assert paragraph.source.kind == "docx"
    table_rows = [b for b in extracted.blocks if b.kind == BlockKind.TABLE_ROW]
    assert any("Rollback" in b.text for b in table_rows)
    assert any(b.kind == BlockKind.HEADER and "Confidential" in b.text for b in extracted.blocks)
