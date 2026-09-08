"""One finding, four formats: content that is present but not visible.

Each extractor answers it in the same words so a caller need not know which fields to
compare per format. What a format has no notion of is left out rather than invented.
"""

from __future__ import annotations

from collections.abc import Callable, Sequence
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn
from openpyxl import Workbook
from pptx import Presentation

from doc_agent.adapters.extractors.docx import DocxExtractor
from doc_agent.adapters.extractors.pdf import PdfExtractor
from doc_agent.adapters.extractors.pptx import PptxExtractor
from doc_agent.adapters.extractors.xlsx import XlsxExtractor


def _tracked_document(path: Path) -> Path:
    """A document with an unaccepted insertion, an unaccepted deletion, and hidden text.

    python-docx models paragraphs, not revision marks, so these are written as the XML
    Word itself writes.
    """

    document = Document()
    document.add_paragraph("Clause 1: agreed text.")
    paragraph = document.add_paragraph("Clause 2: ")
    paragraph.add_run("kept text")

    insertion = paragraph._p.makeelement(qn("w:ins"), {qn("w:id"): "1", qn("w:author"): "Reviewer"})
    run = paragraph._p.makeelement(qn("w:r"), {})
    text = paragraph._p.makeelement(qn("w:t"), {})
    text.text = " INSERTED"
    run.append(text)
    insertion.append(run)
    paragraph._p.append(insertion)

    deletion = paragraph._p.makeelement(qn("w:del"), {qn("w:id"): "2", qn("w:author"): "Reviewer"})
    deleted_run = paragraph._p.makeelement(qn("w:r"), {})
    deleted_text = paragraph._p.makeelement(qn("w:delText"), {})
    deleted_text.text = " REMOVED CLAUSE"
    deleted_run.append(deleted_text)
    deletion.append(deleted_run)
    paragraph._p.append(deletion)

    hidden = document.add_paragraph().add_run("internal note")
    run_properties = hidden._r.get_or_add_rPr()
    run_properties.append(run_properties.makeelement(qn("w:vanish"), {}))

    field_paragraph = document.add_paragraph()
    anchor = field_paragraph.add_run()
    field = anchor._r.makeelement(qn("w:fldSimple"), {qn("w:instr"): " DATE "})
    field_run = anchor._r.makeelement(qn("w:r"), {})
    field_text = anchor._r.makeelement(qn("w:t"), {})
    field_text.text = "01/03/2026"
    field_run.append(field_text)
    field.append(field_run)
    field_paragraph._p.append(field)

    document.save(path)
    return path


def test_a_docx_reports_revisions_that_have_not_been_accepted(tmp_path: Path) -> None:
    """The finding with the largest consequence: the file is not what it appears to be."""

    document = DocxExtractor().extract(_tracked_document(tmp_path / "tracked.docx"))
    revisions = document.metadata["revisions"]

    assert (revisions["insertions"], revisions["deletions"]) == (1, 1)
    assert revisions["authors"] == ["Reviewer"]
    # Deleted text is still in the file, which is the point of reporting it.
    assert revisions["deleted_text"] == [" REMOVED CLAUSE"]
    assert any("still in the file" in line for line in document.metadata["withheld_content"])


def test_a_docx_reports_hidden_text_and_field_results(tmp_path: Path) -> None:
    document = DocxExtractor().extract(_tracked_document(tmp_path / "tracked.docx"))

    assert document.metadata["hidden_text_runs"] == 1
    assert any("hidden" in line for line in document.metadata["withheld_content"])

    # A field is the document counterpart of a formula: a saved result, never recomputed.
    field = document.metadata["fields"][0]
    assert field["instruction"] == "DATE"
    assert field["result"] == "01/03/2026"
    assert field["value_state"] == "cached"


def test_a_clean_docx_withholds_nothing(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Plain text.")
    path = tmp_path / "clean.docx"
    document.save(path)

    extracted = DocxExtractor().extract(path)

    assert extracted.metadata["withheld_content"] == []
    assert extracted.metadata["revisions"]["deletions"] == 0
    assert extracted.metadata["protection"] == {"enabled": False}


def test_a_pptx_reports_hidden_slides(tmp_path: Path) -> None:
    """A slide set never to show is the deck's hidden worksheet."""

    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[1])
    hidden = presentation.slides.add_slide(presentation.slide_layouts[1])
    hidden._element.set("show", "0")
    path = tmp_path / "deck.pptx"
    presentation.save(path)

    extracted = PptxExtractor().extract(path)

    assert extracted.metadata["hidden_slides"] == [2]
    assert extracted.metadata["slide_count"] == 2
    assert any("hidden" in line for line in extracted.metadata["withheld_content"])
    per_slide = {c.ordinal: c.metadata["hidden"] for c in extracted.containers}
    assert per_slide == {1: False, 2: True}


def test_an_xlsx_reports_hidden_sheets_in_the_same_words(tmp_path: Path) -> None:
    workbook = Workbook()
    visible = workbook.active
    assert visible is not None
    visible.append(["a", 1])
    hidden = workbook.create_sheet("Internal")
    hidden.append(["secret", 2])
    hidden.sheet_state = "hidden"
    path = tmp_path / "book.xlsx"
    workbook.save(path)

    extracted = XlsxExtractor().extract(path)

    assert any("Internal" in line for line in extracted.metadata["withheld_content"])


def test_a_pdf_reports_pages_it_could_not_read(
    tmp_path: Path, write_pdf: Callable[[Sequence[dict[str, Any]]], bytes]
) -> None:
    """A scan is absent from the store, not empty in the source."""

    path = tmp_path / "doc.pdf"
    path.write_bytes(write_pdf([{"lines": [(72, 700, 12, "Clause one")]}, {"lines": []}]))

    extracted = PdfExtractor().extract(path)

    assert extracted.metadata["pages_without_text"] == [2]
    assert extracted.metadata["encrypted"] is False
    assert any("No OCR" in line for line in extracted.metadata["withheld_content"])


@pytest.mark.parametrize("key", ["calculation", "revisions", "hidden_slides"])
def test_a_pdf_does_not_borrow_notions_it_has_no_counterpart_for(
    tmp_path: Path, key: str, write_pdf: Callable[[Sequence[dict[str, Any]]], bytes]
) -> None:
    """A PDF has no formulas, no tracked changes and no slides; none are faked for it."""

    path = tmp_path / "doc.pdf"
    path.write_bytes(write_pdf([{"lines": [(72, 700, 12, "Clause one")]}]))

    assert key not in PdfExtractor().extract(path).metadata
