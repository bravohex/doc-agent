"""Workbook facts an audit needs that no cell value can report.

What a cell was allowed to contain, which rules watch it, what names its formulas use,
and whether the workbook recalculates at all are all properties of the file rather than
of any value in it.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from doc_agent.adapters.extractors.xlsx import XlsxExtractor
from doc_agent.domain.models import ExtractedDocument


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "MOG"
    ws.append(["ID", "Status", "Amount", "Total"])
    for index in range(1, 4):
        row = index + 1
        ws.append([f"MOG-{index:03d}", "Open", 100 * index, f"=C{row}*1.1"])

    status = DataValidation(
        type="list",
        formula1='"Open,Closed,Blocked"',
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Invalid status",
        error="Pick from the list",
    )
    status.add("B2:B100")
    ws.add_data_validation(status)
    positive = DataValidation(
        type="decimal", operator="greaterThan", formula1="0", allow_blank=True
    )
    positive.add("C2:C100")
    ws.add_data_validation(positive)

    ws.conditional_formatting.add("C2:C100", CellIsRule(operator="lessThan", formula=["0"]))
    ws.add_table(Table(displayName="FitGap", ref="A1:D4"))

    wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="MOG!$C$1")
    rates = wb.create_sheet("Rates")
    rates.append(["rate", 0.1])
    rates.defined_names["LocalRate"] = DefinedName("LocalRate", attr_text="Rates!$B$1")

    wb.calculation.calcMode = "manual"
    path = tmp_path / "audit.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def extracted(workbook: Path) -> ExtractedDocument:
    return XlsxExtractor().extract(workbook)


def _sheet(document: ExtractedDocument, title: str) -> dict:
    return next(c.metadata for c in document.containers if c.title == title)


def test_validation_rules_are_recorded_with_the_ranges_they_cover(
    extracted: ExtractedDocument,
) -> None:
    rules = _sheet(extracted, "MOG")["validations"]

    assert [rule["type"] for rule in rules] == ["list", "decimal"]
    listed = rules[0]
    assert listed["ranges"] == ["B2:B100"]
    assert listed["formula1"] == '"Open,Closed,Blocked"'
    assert listed["allow_blank"] is False
    # The message the author wrote is part of the rule's intent.
    assert listed["error_title"] == "Invalid status"
    assert listed["error_message"] == "Pick from the list"

    numeric = rules[1]
    assert (numeric["operator"], numeric["formula1"]) == ("greaterThan", "0")
    assert numeric["ranges"] == ["C2:C100"]


def test_conditional_rules_are_recorded_as_conditions_not_colours(
    extracted: ExtractedDocument,
) -> None:
    formats = _sheet(extracted, "MOG")["conditional_formats"]

    assert len(formats) == 1
    assert formats[0]["ranges"] == ["C2:C100"]
    assert formats[0]["type"] == "cellIs"
    assert formats[0]["operator"] == "lessThan"
    assert formats[0]["formula"] == ["0"]


def test_defined_names_are_recorded_at_the_scope_that_owns_them(
    extracted: ExtractedDocument,
) -> None:
    workbook_names = extracted.metadata["defined_names"]
    assert [(n["name"], n["refers_to"], n["scope"]) for n in workbook_names] == [
        ("TaxRate", "MOG!$C$1", "workbook")
    ]

    sheet_names = _sheet(extracted, "Rates")["defined_names"]
    assert [(n["name"], n["scope"]) for n in sheet_names] == [("LocalRate", "Rates")]
    # A workbook-scoped name is not repeated onto every sheet.
    assert _sheet(extracted, "MOG")["defined_names"] == []


def test_manual_calculation_is_recorded_because_it_changes_how_values_read(
    extracted: ExtractedDocument,
) -> None:
    calculation = extracted.metadata["calculation"]

    assert calculation["mode"] == "manual"
    assert calculation["automatic"] is False


def test_an_automatic_workbook_says_so(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["a", 1])
    path = tmp_path / "plain.xlsx"
    wb.save(path)

    calculation = XlsxExtractor().extract(path).metadata["calculation"]

    assert calculation["automatic"] is True
    assert calculation["mode"] != "manual"


def test_a_sheet_without_rules_records_an_empty_list_not_nothing(
    extracted: ExtractedDocument,
) -> None:
    """An empty list is the finding "no rules"; a missing key would mean "unknown"."""

    rates = _sheet(extracted, "Rates")

    assert rates["validations"] == []
    assert rates["conditional_formats"] == []
