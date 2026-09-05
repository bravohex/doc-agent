"""Repeated source content must stay ingestible instead of breaking persistence."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from pptx import Presentation

from doc_agent.bootstrap import build_container


def test_docx_with_repeated_paragraphs_and_table_rows_ingests(tmp_path: Path) -> None:
    source = tmp_path / "repeated.docx"
    document = Document()
    document.add_heading("Scope", level=1)
    document.add_paragraph("N/A")
    document.add_paragraph("N/A")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "X"
    table.cell(0, 1).text = "1"
    table.cell(1, 0).text = "X"
    table.cell(1, 1).text = "2"
    document.save(source)

    app = build_container(tmp_path / "home")
    project = app.projects.create("Repeats")
    result = app.ingest.execute(project.id, source)

    assert result.status == "created"
    blocks = app.repository.current_blocks(result.document_id)
    assert len({block.stable_key for block in blocks}) == len(blocks)


def test_pptx_slides_sharing_a_title_keep_distinct_block_keys(tmp_path: Path) -> None:
    source = tmp_path / "agenda.pptx"
    presentation = Presentation()
    for body in ("first half", "second half"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Agenda"
        slide.placeholders[1].text = body
    presentation.save(source)

    app = build_container(tmp_path / "home")
    project = app.projects.create("Deck")
    result = app.ingest.execute(project.id, source)

    assert result.status == "created"
    blocks = app.repository.current_blocks(result.document_id)
    assert len({block.stable_key for block in blocks}) == len(blocks)
    assert {"first half", "second half"} <= {block.text for block in blocks}
