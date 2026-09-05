from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from doc_agent.adapters.extractors.pptx import PptxExtractor
from doc_agent.bootstrap import build_container
from doc_agent.domain.models import BlockKind


def test_pptx_preserves_slide_text_tables_and_notes(tmp_path: Path) -> None:
    path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Target Architecture"
    box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(1))
    box.text = "Shopify → iPaaS → KUROSHIO"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1.5)).table
    table.cell(0, 0).text = "IF"
    table.cell(0, 1).text = "System"
    table.cell(1, 0).text = "IF-01"
    table.cell(1, 1).text = "KUROSHIO"
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Important integration notes"
    prs.save(path)

    extracted = PptxExtractor().extract(path)
    assert any(b.kind == BlockKind.TEXT_BOX and "Shopify" in b.text for b in extracted.blocks)
    assert any(b.kind == BlockKind.TABLE_ROW and "IF-01" in b.text for b in extracted.blocks)
    assert any(b.kind == BlockKind.NOTE and "integration notes" in b.text for b in extracted.blocks)
    assert extracted.containers[0].title == "Target Architecture"


def test_pptx_shape_without_explicit_geometry_is_extracted(tmp_path: Path) -> None:
    source = tmp_path / "inherited.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "Inherited geometry"
    xfrm = box._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
    box._element.spPr.remove(xfrm)
    presentation.save(source)

    app = build_container(tmp_path / "home")
    project = app.projects.create("Inherited")
    result = app.ingest.execute(project.id, source)

    assert result.status == "created"
    block = next(
        b for b in app.repository.current_blocks(result.document_id) if "Inherited" in b.text
    )
    assert block.presentation["left"] is None
