"""Selective formatting: what hides content, and what a reviewer reads meaning into.

Two rules shape this. Sheet-wide layout costs one record per sheet, so all of it is
kept. Cell styling is kept only where it deviates from a plain cell, or the map would be
as large as the sheet while saying nothing.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.styles.colors import Color

from doc_agent.adapters.extractors.xlsx import XlsxExtractor, _as_ranges
from doc_agent.domain.models import ExtractedDocument

_PLAIN = ("coordinate", "number_format", "merged_range", "hidden_column")


@pytest.fixture
def formatted(tmp_path: Path) -> ExtractedDocument:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "MOG"
    ws.append(["ID", "Status", "Amount", "Note"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(["MOG-001", "Open", 100, "ok"])
    ws.append(["MOG-002", "Void", 0, "cancelled"])
    for cell in ws[3]:
        cell.font = Font(strike=True, color="FFFF0000")
    ws.append(["MOG-003", "Open", 300, "check"])
    ws["C4"].fill = PatternFill("solid", start_color="FFFFFF00")
    ws["D4"].protection = Protection(locked=False)

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:D4"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["D"].hidden = True
    ws.row_dimensions[3].hidden = True
    ws.protection.sheet = True

    path = tmp_path / "fmt.xlsx"
    wb.save(path)
    return XlsxExtractor().extract(path)


def _layout(document: ExtractedDocument) -> dict[str, Any]:
    return next(c.metadata["layout"] for c in document.containers if c.title == "MOG")


def _style(document: ExtractedDocument, coordinate: str) -> dict[str, Any]:
    for block in document.blocks:
        for layout in block.presentation.get("cells", []):
            if layout.get("coordinate") == coordinate:
                return {key: value for key, value in layout.items() if key not in _PLAIN}
    raise AssertionError(f"no cell at {coordinate}")


@pytest.mark.parametrize(
    ("numbers", "expected"),
    [
        ([], []),
        ([3], [[3, 3]]),
        ([3, 4, 5], [[3, 5]]),
        ([3, 4, 5, 9], [[3, 5], [9, 9]]),
        ([1, 3, 5], [[1, 1], [3, 3], [5, 5]]),
    ],
)
def test_hidden_rows_compress_into_ranges(numbers: list[int], expected: list[list[int]]) -> None:
    """A sheet with thousands of folded rows must stay a short answer."""

    assert _as_ranges(numbers) == expected


def test_sheet_layout_records_what_can_hide_content(formatted: ExtractedDocument) -> None:
    layout = _layout(formatted)

    assert layout["freeze_panes"] == "B2"
    assert layout["auto_filter"] == "A1:D4"
    assert layout["protected"] is True
    assert layout["hidden_rows"] == [[3, 3]]
    columns = {entry["column"]: entry for entry in layout["columns"]}
    assert columns["D"]["hidden"] is True
    assert columns["A"]["width"] == 18.0
    # Columns that were never sized or hidden are not listed.
    assert "B" not in columns


def test_only_deviating_cell_styling_is_recorded(formatted: ExtractedDocument) -> None:
    assert _style(formatted, "A1") == {"bold": True}
    assert _style(formatted, "A3") == {"strikethrough": True, "font_color": "FFFF0000"}
    assert _style(formatted, "C4") == {"fill_color": "FFFFFF00"}
    # Locked is the default, so only the editable cell records anything.
    assert _style(formatted, "D4") == {"locked": False}
    # A plain cell adds nothing at all.
    assert _style(formatted, "A2") == {}


def test_an_automatic_theme_colour_is_not_recorded_as_a_choice(tmp_path: Path) -> None:
    """Themes 0 and 1 are the automatic colours nearly every cell carries.

    Recording them would make the sparse style map dense again, and openpyxl reports a
    theme colour by returning its validation *message* from ``rgb`` -- so a naive read
    stores the sentence "Values must be of type <class 'str'>" where a colour belongs.
    """

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["plain"])
    path = tmp_path / "theme.xlsx"
    wb.save(path)

    document = XlsxExtractor().extract(path)
    style = _style(document, "A1")

    assert "font_color" not in style
    assert "Values must be" not in str(style)


def test_a_deliberate_theme_colour_is_reported_as_its_theme_index(tmp_path: Path) -> None:
    """The workbook palette is not resolved, so naming a concrete colour would be a guess.

    This is also where openpyxl's placeholder bites hardest: for a theme colour, ``rgb``
    returns the string "Values must be of type <class 'str'>".
    """

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["accent"])
    ws["A1"].font = Font(color=Color(theme=4, tint=0.0))
    path = tmp_path / "accent.xlsx"
    wb.save(path)

    style = _style(XlsxExtractor().extract(path), "A1")

    assert style["font_color"] == "theme:4"


def test_an_explicit_rgb_colour_survives_the_theme_guard(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["green"])
    ws["A1"].font = Font(color="FF00FF00")
    path = tmp_path / "rgb.xlsx"
    wb.save(path)

    assert _style(XlsxExtractor().extract(path), "A1")["font_color"] == "FF00FF00"
