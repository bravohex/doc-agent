from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlsxImage
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from doc_agent.adapters.extractors.docx import DocxExtractor
from doc_agent.adapters.extractors.pptx import PptxExtractor
from doc_agent.adapters.extractors.xlsx import XlsxExtractor


def _png(path: Path) -> None:
    Image.new("RGB", (8, 8), "white").save(path)


def test_all_office_extractors_register_embedded_images(tmp_path: Path) -> None:
    image = tmp_path / "image.png"
    _png(image)

    xlsx = tmp_path / "image.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "visual"
    ws.add_image(XlsxImage(image), "C3")
    wb.save(xlsx)
    x = XlsxExtractor().extract(xlsx)
    assert len(x.visuals) == 1
    assert x.visuals[0].media_type == "image/png"

    docx = tmp_path / "image.docx"
    doc = Document()
    doc.add_picture(str(image))
    doc.save(docx)
    d = DocxExtractor().extract(docx)
    assert len(d.visuals) == 1
    assert d.visuals[0].media_type == "image/png"

    pptx = tmp_path / "image.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.add_picture(str(image), Inches(1), Inches(1))
    prs.save(pptx)
    p = PptxExtractor().extract(pptx)
    assert len(p.visuals) == 1
    assert p.visuals[0].media_type == "image/png"
