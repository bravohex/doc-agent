from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment

from doc_agent.adapters.extractors.xlsx import XlsxExtractor
from doc_agent.domain.models import BlockKind


def test_xlsx_preserves_formula_format_comment_merge_and_hidden_metadata(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "MOG"
    ws.append(["ID", "Function", "Rate", "Formula"])
    ws.append(["MOG-001", "PayPay", 0.125, "=3*4"])
    ws["C2"].number_format = "0.0%"
    ws["B2"].comment = Comment("business note", "OLM")
    ws.merge_cells("A4:C4")
    ws["A4"] = "Merged heading"
    ws.row_dimensions[2].hidden = True
    ws.column_dimensions["D"].hidden = True
    wb.save(path)

    doc = XlsxExtractor().extract(path)
    rows = [b for b in doc.blocks if b.kind == BlockKind.TABLE_ROW]
    assert any("PayPay" in b.text for b in rows)
    row = next(b for b in rows if "MOG-001" in b.text)
    # Semantic cells and their position/style entries share one index.
    cells = row.payload["cells"]
    layout = row.presentation["cells"]
    by_coordinate = {
        str(position["coordinate"]): cell for cell, position in zip(cells, layout, strict=True)
    }
    assert by_coordinate["C2"]["raw_value"] == 0.125
    assert by_coordinate["C2"]["display"] == "12.5%"
    assert by_coordinate["D2"]["formula"] == "=3*4"
    assert by_coordinate["B2"]["comment"] == "business note"
    assert next(p for p in layout if p["coordinate"] == "C2")["number_format"] == "0.0%"
    assert next(p for p in layout if p["coordinate"] == "D2")["hidden_column"] is True
    assert row.presentation["hidden_row"] is True
    merged = next(b for b in rows if "Merged heading" in b.text)
    assert merged.presentation["cells"][0]["merged_range"] == "A4:C4"
    assert len([c for c in merged.payload["cells"] if c["raw_value"] == "Merged heading"]) == 1


def test_xlsx_row_identity_survives_row_insertion(tmp_path: Path) -> None:
    """A physical row move must not change identity when row semantics are unchanged."""

    first_path = tmp_path / "first.xlsx"
    second_path = tmp_path / "second.xlsx"

    first = Workbook()
    first_ws = first.active
    first_ws.title = "MOG"
    first_ws.append(["MOG-001", "PayPay", "A"])
    first.save(first_path)

    second = Workbook()
    second_ws = second.active
    second_ws.title = "MOG"
    second_ws.append(["Inserted heading"])
    second_ws.append(["MOG-001", "PayPay", "A"])
    second.save(second_path)

    extractor = XlsxExtractor()
    first_doc = extractor.extract(first_path)
    second_doc = extractor.extract(second_path)
    first_row = next(block for block in first_doc.blocks if "MOG-001" in block.text)
    second_row = next(block for block in second_doc.blocks if "MOG-001" in block.text)

    assert first_row.stable_key == second_row.stable_key
    assert first_row.source.row == 1
    assert second_row.source.row == 2
